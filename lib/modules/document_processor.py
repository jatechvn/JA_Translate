# lib/modules/document_processor.py
# Python backend helper for layout-preserving document translation (PDF, Excel, PowerPoint, Word)

import os
import sys
import json
import time
import argparse
import configparser
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed

warnings.filterwarnings("ignore")

# ─── Bootstrap Dependencies ──────────────────────────────────────────────────
REQUIRED_LIBS = {
    'openpyxl': 'openpyxl',
    'pptx': 'python-pptx',
    'docx': 'python-docx',
    'fitz': 'pymupdf',
    'requests': 'requests'
}

def bootstrap_dependencies():
    for module_name, pip_name in REQUIRED_LIBS.items():
        try:
            __import__(module_name)
        except ImportError:
            print(f"Installing missing dependency: {pip_name}...", flush=True)
            import subprocess
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", pip_name])
            except Exception as e:
                print(f"Failed to install {pip_name}: {e}", flush=True)
                sys.exit(1)

bootstrap_dependencies()

# Now import the dependencies safely
import requests
import openpyxl
import fitz  # PyMuPDF
from pptx import Presentation
from docx import Document

# ─── Config Parser ───────────────────────────────────────────────────────────
def load_config():
    config = configparser.ConfigParser()
    # Look for config.ini in cwd, and also check project root relative to this script
    possible_paths = [
        'config.ini',
        os.path.join(os.path.dirname(__file__), '../../config.ini'),
        os.path.join(os.path.dirname(__file__), '../config.ini'),
    ]
    config_path = None
    for path in possible_paths:
        if os.path.exists(path):
            config_path = path
            break
            
    if not config_path:
        raise Exception("Configuration file config.ini not found.")
        
    config.read(config_path)
    return config

# ─── Translation Cache ────────────────────────────────────────────────────────
_cache = {}
_cache_path = 'translation_cache.json'

def load_cache():
    global _cache, _cache_path
    possible_paths = [
        'translation_cache.json',
        os.path.join(os.path.dirname(__file__), '../../translation_cache.json'),
        os.path.join(os.path.dirname(__file__), '../translation_cache.json'),
    ]
    for path in possible_paths:
        if os.path.exists(path):
            _cache_path = path
            break
    if os.path.exists(_cache_path):
        try:
            with open(_cache_path, 'r', encoding='utf-8') as f:
                _cache = json.load(f)
        except Exception:
            _cache = {}

def save_cache():
    try:
        with open(_cache_path, 'w', encoding='utf-8') as f:
            json.dump(_cache, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

def get_cached_translation(text, src_lang, tgt_lang):
    key = f"{src_lang}->{tgt_lang}:{text}"
    return _cache.get(key)

def set_cached_translation(text, src_lang, tgt_lang, translated):
    key = f"{src_lang}->{tgt_lang}:{text}"
    _cache[key] = translated

# ─── Translation API Client ──────────────────────────────────────────────────
def translate_text(text, src_lang, tgt_lang, api_key, api_url, model, proxy_settings=None):
    if not text or not text.strip():
        return text
        
    cached = get_cached_translation(text, src_lang, tgt_lang)
    if cached:
        return cached
        
    lang_map = {
        'VN': 'Vietnamese',
        'ENG': 'English',
        'CN': 'Chinese (Simplified)',
    }
    
    source_name = lang_map.get(src_lang, 'Auto-detect')
    target_name = lang_map.get(tgt_lang, 'Vietnamese')
    
    system_prompt = f"You are a professional translator. Translate from {source_name} to {target_name}. Maintain meaning and formatting. Only output the translation result."
    
    headers = {
        'Authorization': f'Bearer {api_key}',
        'Content-Type': 'application/json',
    }
    
    payload = {
        'model': model,
        'messages': [
            {'role': 'system', 'content': system_prompt},
            {'role': 'user', 'content': text}
        ],
        'temperature': 0.1,
    }
    
    proxies = None
    if proxy_settings and proxy_settings.get('enabled') == 'true':
        host = proxy_settings.get('host')
        port = proxy_settings.get('port')
        user = proxy_settings.get('user')
        pwd = proxy_settings.get('pass')
        if host and port:
            if user and pwd:
                proxy_url = f"http://{user}:{pwd}@{host}:{port}"
            else:
                proxy_url = f"http://{host}:{port}"
            proxies = {
                'http': proxy_url,
                'https': proxy_url
            }
            
    # Simple retry block with exponential backoff and jitter for 429 / Too Many Requests
    import random
    max_retries = 5
    for attempt in range(max_retries):
        try:
            r = requests.post(f"{api_url}/chat/completions", headers=headers, json=payload, proxies=proxies, timeout=45)
            
            if r.status_code == 429:
                retry_after = r.headers.get("Retry-After")
                wait_time = 0
                if retry_after:
                    try:
                        wait_time = float(retry_after)
                    except ValueError:
                        pass
                if not wait_time:
                    wait_time = (2 ** attempt) + random.uniform(0.5, 1.5)
                
                time.sleep(wait_time)
                continue
                
            r.raise_for_status()
            res = r.json()
            translated = res['choices'][0]['message']['content'].strip()
            set_cached_translation(text, src_lang, tgt_lang, translated)
            return translated
        except Exception as e:
            if attempt == max_retries - 1:
                raise e
            wait_time = (2 ** attempt) + random.uniform(0.5, 1.5)
            time.sleep(wait_time)

# ─── Parallel Translation Helper ─────────────────────────────────────────────
def translate_items(items, get_text_fn, set_text_fn, src_lang, tgt_lang, api_key, api_url, model, proxy, status_template):
    # Gather all texts and map items to their texts
    item_texts = []
    unique_texts = set()
    for item in items:
        text = get_text_fn(item)
        item_texts.append((item, text))
        if text and text.strip():
            unique_texts.add(text)
            
    total_unique = len(unique_texts)
    if total_unique == 0:
        return
        
    # Translate unique texts in parallel, checking cache first
    translations = {}
    uncached_texts = set()
    cached_count = 0
    for text in unique_texts:
        cached = get_cached_translation(text, src_lang, tgt_lang)
        if cached:
            translations[text] = cached
            cached_count += 1
        else:
            uncached_texts.add(text)
            
    total_uncached = len(uncached_texts)
    
    if total_uncached > 0:
        translated_count = 0
        # Concurrency limit: 3 concurrent workers is safe for rate limits
        max_workers = min(3, total_uncached)
        
        def translate_single(text):
            try:
                val = translate_text(text, src_lang, tgt_lang, api_key, api_url, model, proxy)
                return text, val, None
            except Exception as e:
                return text, None, str(e)

        initial_status = f"Found {cached_count} cached translations. Translating remaining {total_uncached} texts..."
        print(json.dumps({"progress": 0.05, "status": initial_status}), flush=True)

        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {executor.submit(translate_single, text): text for text in uncached_texts}
            
            for future in as_completed(futures):
                text, val, err = future.result()
                if err:
                    raise Exception(f"Translation failed: {err}")
                
                translations[text] = val
                translated_count += 1
                
                progress = 0.05 + 0.90 * ((cached_count + translated_count) / total_unique)
                status_str = status_template.format(
                    current=cached_count + translated_count, 
                    total=total_unique
                ) + f" ({cached_count} from cache)"
                print(json.dumps({
                    "progress": progress,
                    "status": status_str
                }), flush=True)
    else:
        # 100% Cache hit!
        progress = 0.95
        status_str = f"All {total_unique} texts loaded from local cache!"
        print(json.dumps({
            "progress": progress,
            "status": status_str
        }), flush=True)
            
    # Apply translations back to items
    for item, original_text in item_texts:
        if original_text and original_text.strip():
            translated_val = translations.get(original_text, original_text)
            set_text_fn(item, translated_val)

# ─── PDF Vietnamese Font Optimizer ──────────────────────────────────────────
def get_vietnamese_font():
    # Common Windows paths
    paths = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        r"C:\Windows\Fonts\times.ttf",
    ]
    # Common macOS paths
    paths += [
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    # Common Linux paths
    paths += [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    
    for path in paths:
        if os.path.exists(path):
            name = os.path.basename(path).split('.')[0].lower()
            return path, name
            
    return None, "helv"

# ─── PDF Translator (Layout Preserving) ──────────────────────────────────────
def process_pdf(input_path, output_path, src_lang, tgt_lang, api_key, api_url, model, proxy):
    doc = fitz.open(input_path)
    
    print(json.dumps({"progress": 0.05, "status": "Analyzing PDF layout..."}), flush=True)
    
    target_blocks = []
    for page_idx, page in enumerate(doc):
        # b is (x0, y0, x1, y1, text, block_no, block_type)
        blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
        for b in blocks:
            target_blocks.append({
                'page': page,
                'rect': fitz.Rect(b[0], b[1], b[2], b[3]),
                'text': b[4]
            })
            
    if not target_blocks:
        doc.close()
        raise Exception("No text content found in PDF.")
        
    translate_items(
        items=target_blocks,
        get_text_fn=lambda b: b['text'],
        set_text_fn=lambda b, val: b.update({'translated_text': val}),
        src_lang=src_lang,
        tgt_lang=tgt_lang,
        api_key=api_key,
        api_url=api_url,
        model=model,
        proxy=proxy,
        status_template="Translating PDF text block {current} of {total}..."
    )
    
    font_file, font_name = get_vietnamese_font()
    
    print(json.dumps({"progress": 0.95, "status": "Applying PDF modifications..."}), flush=True)
    for b in target_blocks:
        page = b['page']
        rect = b['rect']
        translated = b.get('translated_text', b['text'])
        
        # Apply redaction to clear the original text
        page.add_redact_annot(rect)
        page.apply_redactions(fill=(1, 1, 1))
        
        # Insert translated text into the textbox area
        page.insert_textbox(rect, translated, fontfile=font_file, fontname=font_name, fontsize=10)
        
    doc.save(output_path)
    doc.close()

# ─── Excel Translator ────────────────────────────────────────────────────────
def process_excel(input_path, output_path, src_lang, tgt_lang, api_key, api_url, model, proxy):
    wb = openpyxl.load_workbook(input_path)
    sheets = wb.sheetnames
    
    print(json.dumps({"progress": 0.05, "status": "Scanning Excel sheets..."}), flush=True)
    
    target_cells = []
    for sheet_name in sheets:
        ws = wb[sheet_name]
        for cell in ws._cells.values():
            if cell.value and isinstance(cell.value, str) and not str(cell.value).startswith('='):
                target_cells.append(cell)
                    
    if not target_cells:
        wb.close()
        raise Exception("No text cells found in Excel sheet.")
        
    def set_cell_value(c, val):
        c.value = val

    translate_items(
        items=target_cells,
        get_text_fn=lambda c: c.value,
        set_text_fn=set_cell_value,
        src_lang=src_lang,
        tgt_lang=tgt_lang,
        api_key=api_key,
        api_url=api_url,
        model=model,
        proxy=proxy,
        status_template="Translating Excel cells {current} of {total}..."
    )
        
    wb.save(output_path)
    wb.close()

# ─── PowerPoint Translator (Layout Preserving) ───────────────────────────────
def process_pptx(input_path, output_path, src_lang, tgt_lang, api_key, api_url, model, proxy):
    prs = Presentation(input_path)
    
    print(json.dumps({"progress": 0.05, "status": "Scanning slides..."}), flush=True)
    
    target_paragraphs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text and paragraph.text.strip():
                        target_paragraphs.append(paragraph)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            if paragraph.text and paragraph.text.strip():
                                target_paragraphs.append(paragraph)
                                
    if not target_paragraphs:
        raise Exception("No text frames found in PowerPoint.")
        
    def set_paragraph_text(p, val):
        if len(p.runs) > 0:
            p.runs[0].text = val
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.text = val

    translate_items(
        items=target_paragraphs,
        get_text_fn=lambda p: p.text,
        set_text_fn=set_paragraph_text,
        src_lang=src_lang,
        tgt_lang=tgt_lang,
        api_key=api_key,
        api_url=api_url,
        model=model,
        proxy=proxy,
        status_template="Translating PowerPoint texts {current} of {total}..."
    )
        
    prs.save(output_path)

# ─── Word Translator (Layout Preserving) ─────────────────────────────────────
def process_docx(input_path, output_path, src_lang, tgt_lang, api_key, api_url, model, proxy):
    doc = Document(input_path)
    
    print(json.dumps({"progress": 0.05, "status": "Scanning Word paragraphs..."}), flush=True)
    
    target_paragraphs = []
    for paragraph in doc.paragraphs:
        if paragraph.text and paragraph.text.strip():
            target_paragraphs.append(paragraph)
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text and paragraph.text.strip():
                        target_paragraphs.append(paragraph)
                        
    if not target_paragraphs:
        raise Exception("No text paragraphs found in Word document.")
        
    def set_paragraph_text(p, val):
        if len(p.runs) > 0:
            p.runs[0].text = val
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.text = val

    translate_items(
        items=target_paragraphs,
        get_text_fn=lambda p: p.text,
        set_text_fn=set_paragraph_text,
        src_lang=src_lang,
        tgt_lang=tgt_lang,
        api_key=api_key,
        api_url=api_url,
        model=model,
        proxy=proxy,
        status_template="Translating Word text {current} of {total}..."
    )
        
    doc.save(output_path)

# ─── Main Execution ──────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Document Translation Script")
    parser.add_argument("--input", required=True, help="Input file path")
    parser.add_argument("--output", required=False, help="Output file path")
    parser.add_argument("--src", required=True, help="Source language")
    parser.add_argument("--tgt", required=True, help="Target language")
    parser.add_argument("--cache", required=False, help="Path to translation cache JSON file")
    args = parser.parse_args()
    
    try:
        ext = args.input.split('.')[-1].lower()
        
        # Stat extraction only
        if args.src == 'stats':
            text = ""
            if ext == 'xlsx':
                wb = openpyxl.load_workbook(args.input, data_only=True)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    for cell in ws._cells.values():
                        val = cell.value
                        if val and isinstance(val, str) and not str(val).startswith('='):
                            text += val + " "
                wb.close()
            elif ext == 'pptx':
                prs = Presentation(args.input)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            text += shape.text_frame.text + " "
                        if shape.has_table:
                            for row in shape.table.rows:
                                for cell in row.cells:
                                    text += cell.text_frame.text + " "
            elif ext == 'docx':
                doc = Document(args.input)
                for p in doc.paragraphs:
                    text += p.text + " "
                for t in doc.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            text += cell.text + " "
            elif ext == 'pdf':
                doc = fitz.open(args.input)
                for page in doc:
                    text += page.get_text() + " "
                doc.close()
            else:
                raise Exception(f"Unsupported stats file type: .{ext}")
                
            char_count = len(text)
            word_count = len([w for w in text.split() if w.strip()])
            print(json.dumps({"char_count": char_count, "word_count": word_count}), flush=True)
            sys.exit(0)

        # Full translation execution
        if not args.output:
            raise Exception("Output file path is required for translation.")

        # Load API keys and Proxy
        config = load_config()
        if args.cache:
            global _cache_path
            _cache_path = args.cache
        load_cache()
        
        active_provider = config.get('SETTINGS', 'active_provider', fallback='cloud').lower()
        if active_provider in ('local', 'local_ai'):
            api_key = ''
            llama_port = config.get('LOCAL_AI', 'llama_port', fallback='8080')
            api_url = config.get('LOCAL_AI', 'endpoint', fallback=f'http://127.0.0.1:{llama_port}/v1').rstrip('/')
            model = config.get('LOCAL_AI', 'gguf_model', fallback='qwen2.5-1.5b-instruct-q4_k_m.gguf')
        else:
            api_key = config.get('NVIDIA', 'api_key')
            api_url = config.get('NVIDIA', 'api_base', fallback='https://integrate.api.nvidia.com/v1')
            model = config.get('NVIDIA', 'model', fallback='google/gemma-4-31b-it')
        
        proxy = {
            'enabled': config.get('PROXY', 'enabled', fallback='false').lower(),
            'host': config.get('PROXY', 'host', fallback=''),
            'port': config.get('PROXY', 'port', fallback=''),
            'user': config.get('PROXY', 'user', fallback=''),
            'pass': config.get('PROXY', 'pass', fallback=''),
        }
        
        if ext == 'pdf':
            process_pdf(args.input, args.output, args.src, args.tgt, api_key, api_url, model, proxy)
        elif ext == 'xlsx':
            process_excel(args.input, args.output, args.src, args.tgt, api_key, api_url, model, proxy)
        elif ext == 'pptx':
            process_pptx(args.input, args.output, args.src, args.tgt, api_key, api_url, model, proxy)
        elif ext == 'docx':
            process_docx(args.input, args.output, args.src, args.tgt, api_key, api_url, model, proxy)
        else:
            raise Exception(f"Unsupported file type: .{ext}")
            
        save_cache()
        print(json.dumps({"progress": 1.0, "status": "complete"}), flush=True)
        sys.exit(0)
        
    except Exception as e:
        print(json.dumps({"progress": 0.0, "status": "error", "error": str(e)}), flush=True)
        sys.exit(1)

if __name__ == '__main__':
    main()
